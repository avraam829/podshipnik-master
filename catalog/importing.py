"""Safe, transactional import of the catalog Excel template.

The importer only creates products or updates products with matching stable
codes.  It never treats an omitted row as a deletion.
"""
import hashlib
import json
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from django.db import transaction
from .models import Category, Product, Price, Stock, Warehouse
from .services import normalize_designation, product_slug


class ImportValidationError(ValueError):
    pass


EXCEL_HEADERS = (
    'Родительская категория', 'Категория', 'Код', 'Маркировка',
    'Производитель', 'Описание', 'Вес', 'Длина', 'Ширина', 'Высота',
    'Склад', 'Количество', 'Зарезервировано', 'Свободно',
    'Цена Розница', 'Цена Опт',
)
DEFAULT_WAREHOUSE = 'Основной склад'
MAX_DATA_ROWS = 100_000
DECIMAL_FIELDS = {
    6: (10, 3),  # weight
    7: (10, 2), 8: (10, 2), 9: (10, 2),
    11: (10, 0), 12: (10, 0),
    14: (10, 2), 15: (10, 2),
}
MAX_STOCK_AMOUNT = 2_147_483_647


def decimal_cell(value, field, line, max_digits, decimal_places):
    if value in (None, ''):
        return None
    try:
        number = Decimal(str(value).replace('\u00a0', '').replace(' ', '').replace(',', '.'))
        unit = Decimal(1).scaleb(-decimal_places)
        maximum = Decimal(10 ** (max_digits - decimal_places)) - unit
        if (not number.is_finite() or number < 0 or number > maximum
                or number != number.quantize(unit)):
            raise InvalidOperation
        return number
    except (InvalidOperation, ValueError):
        raise ImportValidationError(f'Строка {line}: некорректное поле {field}.')


def read_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet.max_row > MAX_DATA_ROWS + 1 or sheet.max_column != len(EXCEL_HEADERS):
            raise ImportValidationError(
                f'Ожидаются {len(EXCEL_HEADERS)} колонок и не более {MAX_DATA_ROWS} строк данных.'
            )
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator)
        def key(value):
            return ''.join(c for c in str(value or '').casefold() if c.isalnum())
        if [key(v) for v in header] != [key(v) for v in EXCEL_HEADERS]:
            raise ImportValidationError(
                'Заголовки Excel не соответствуют шаблону. Скачайте шаблон или выгрузите каталог из админки.'
            )
        rows = []
        warnings = []
        for line, values in enumerate(iterator, 2):
            if not any(v is not None for v in values):
                continue
            row = list(values)
            if row[3] is None or not str(row[3]).strip():
                raise ImportValidationError(f'Строка {line}: отсутствует маркировка.')
            for index in (0, 1, 2, 3, 4, 5, 10):
                row[index] = str(row[index]).strip() if row[index] is not None else ''
            row[2] = row[2].strip()
            if (
                len(row[0]) > 255 or len(row[1]) > 255 or len(row[2]) > 100
                or len(row[3]) > 255 or len(row[4]) > 100 or len(row[10]) > 255
            ):
                raise ImportValidationError(
                    f'Строка {line}: слишком длинные категория, склад, маркировка, бренд или код.'
                )
            for index in (6, 7, 8, 9, 11, 12, 14, 15):
                if index in (14, 15) and str(row[index]).startswith(('#VALUE!', '#N/A', '#REF!', '#DIV/0!', '#NUM!')):
                    warnings.append(
                        f'Строка {line}: ошибка Excel {row[index]} в цене; '
                        'цена не меняется (для нового товара — по запросу).'
                    )
                    row[index] = None
                max_digits, decimal_places = DECIMAL_FIELDS[index]
                row[index] = decimal_cell(
                    row[index], EXCEL_HEADERS[index], line, max_digits, decimal_places
                )
            for index in (11, 12):
                if row[index] is not None and row[index] != row[index].to_integral_value():
                    raise ImportValidationError(f'Строка {line}: количество должно быть целым.')
                if row[index] is not None and row[index] > MAX_STOCK_AMOUNT:
                    raise ImportValidationError(
                        f'Строка {line}: количество не может быть больше {MAX_STOCK_AMOUNT}.'
                    )
            if (row[12] or 0) > (row[11] or 0):
                raise ImportValidationError(f'Строка {line}: резерв больше остатка.')
            rows.append(row)
        return rows, warnings
    finally:
        workbook.close()


def _exact_repeat(rows, prefix):
    """Return True only when a no-code file represents the already imported state."""
    expected = {}
    for index, row in enumerate(rows, 1):
        code = row[2] or f'{prefix}{index:06d}'
        item = expected.setdefault(code, {'row': row, 'stocks': {}, 'prices': {}})
        item['stocks'][row[10] or DEFAULT_WAREHOUSE] = (int(row[11] or 0), int(row[12] or 0))
        item['prices'] = {'retail': row[14], 'wholesale': row[15]}

    products = {
        product.code: product
        for product in Product.objects.filter(code__in=expected).select_related('category__parent').prefetch_related(
            'stocks__warehouse', 'prices'
        )
    }
    if set(products) != set(expected):
        return False

    for code, spec in expected.items():
        row = spec['row']
        product = products[code]
        category = product.category
        if (
            product.designation != row[3]
            or product.manufacturer != row[4]
            or product.description != row[5]
            or product.weight != row[6]
            or product.length != row[7]
            or product.width != row[8]
            or product.height != row[9]
            or product.price != (row[14] if row[14] is not None else Decimal(0))
            or (category.name if category else '') != row[1]
            or (category.parent.name if category and category.parent else '') != row[0]
        ):
            return False
        stocks = {
            stock.warehouse.name: (stock.amount, stock.reserved)
            for stock in product.stocks.all()
        }
        prices = {price.price_type: price.price for price in product.prices.all()}
        expected_prices = {kind: value for kind, value in spec['prices'].items() if value is not None}
        if stocks != spec['stocks'] or prices != expected_prices:
            return False
    return True


def _warehouse_plan(rows):
    """Resolve every warehouse without writing, so ambiguity fails in preview too."""
    names = {row[10] or DEFAULT_WAREHOUSE for row in rows}
    by_name = {}
    for warehouse in Warehouse.objects.filter(name__in=names):
        by_name.setdefault(warehouse.name, []).append(warehouse)

    hashed_codes = {
        name: 'WH-' + hashlib.sha256(name.encode()).hexdigest()[:20]
        for name in names
    }
    by_code = {
        warehouse.code: warehouse
        for warehouse in Warehouse.objects.filter(code__in=hashed_codes.values())
    }
    plan = {}
    for name, hashed_code in hashed_codes.items():
        matches = by_name.get(name, [])
        if len(matches) > 1:
            raise ImportValidationError(
                f'Склад «{name}» найден несколько раз. Исправьте дубли складов в админке; данные не изменены.'
            )
        if matches:
            plan[name] = matches[0]
            continue
        collision = by_code.get(hashed_code)
        if collision and collision.name != name:
            raise ImportValidationError(
                f'Код склада для «{name}» уже занят складом «{collision.name}». Данные не изменены.'
            )
        plan[name] = None
    return plan, hashed_codes


def import_xlsx(path, dry_run=True):
    rows, warnings = read_rows(path)
    digest = hashlib.sha256(json.dumps(rows, ensure_ascii=False, default=str).encode()).hexdigest()
    missing_codes = any(not row[2] for row in rows)
    prefix = f'XLS-{digest[:16]}-'
    warehouse_plan, warehouse_codes = _warehouse_plan(rows)
    # A no-code file can be loaded once into an empty catalog. Its exact repeat is a no-op.
    # A changed no-code file cannot safely be mapped onto existing products.
    if missing_codes and Product.objects.exists():
        if _exact_repeat(rows, prefix):
            return {'rows': len(rows), 'created': 0, 'updated': 0, 'unchanged': len(rows), 'dry_run': dry_run, 'warnings': warnings}
        raise ImportValidationError(
            'В файле есть строки без кода, а каталог уже заполнен. Изменённый такой файл нельзя '
            'безопасно сопоставить с товарами. Выгрузите каталог с кодами, перенесите изменения в '
            'эту выгрузку и повторите проверку. Данные не изменены.'
        )
    seen = {}
    seen_stocks = set()
    for row in rows:
        if row[2]:
            signature = row[:10] + row[14:16]
            if row[2] in seen and seen[row[2]] != signature:
                raise ImportValidationError(f'Код {row[2]} повторён с различающимися характеристиками/ценами. Данные не изменены.')
            stock_key = (row[2], row[10] or DEFAULT_WAREHOUSE)
            if stock_key in seen_stocks:
                raise ImportValidationError(f'Повтор товара {row[2]} на одном складе. Данные не изменены.')
            seen_stocks.add(stock_key)
            seen[row[2]] = signature
    existing = set(Product.objects.filter(code__in=[r[2] for r in rows if r[2]]).values_list('code', flat=True))
    result = {'rows': len(rows), 'created': len(set(r[2] for r in rows if r[2]) - existing) + sum(not r[2] for r in rows), 'updated': len(existing), 'dry_run': dry_run, 'warnings': warnings}
    if dry_run:
        return result
    with transaction.atomic():
        # Recheck the no-code condition inside the transaction; this does not claim a database lock.
        if missing_codes and Product.objects.exists():
            raise ImportValidationError('Каталог изменился во время проверки. Повторите предварительную проверку.')
        categories, warehouses = {}, {}
        for warehouse_name, resolved in warehouse_plan.items():
            if resolved is not None:
                warehouses[warehouse_name] = resolved
                continue
            warehouse, _ = Warehouse.objects.get_or_create(
                code=warehouse_codes[warehouse_name], defaults={'name': warehouse_name}
            )
            if warehouse.name != warehouse_name:
                raise ImportValidationError(
                    f'Код склада для «{warehouse_name}» уже занят складом «{warehouse.name}». Данные не изменены.'
                )
            warehouses[warehouse_name] = warehouse
        new_products, stock_specs, price_specs = {}, [], {}
        for i, row in enumerate(rows, 1):
            parent_name, name, code, designation, brand, description = row[:6]
            parent = None
            if parent_name:
                parent = categories.setdefault((parent_name, None), None)
                if parent is None:
                    parent, _ = Category.objects.get_or_create(name=parent_name, parent=None)
                    categories[(parent_name, None)] = parent
            category = None
            if name:
                key = (name, parent.pk if parent else None)
                if key not in categories:
                    categories[key], _ = Category.objects.get_or_create(name=name, parent=parent)
                category = categories[key]
            code = code or f'{prefix}{i:06d}'
            fields = dict(designation=designation, manufacturer=brand, description=description,
                weight=row[6], length=row[7], width=row[8], height=row[9], category=category)
            if code in existing:
                product = Product.objects.get(code=code)
                # An empty/error price means "not supplied", never "delete existing price".
                if row[14] is not None:
                    fields['price'] = row[14]
                for field, value in fields.items():
                    setattr(product, field, value)
                product.save()
            else:
                fields['price'] = row[14] if row[14] is not None else Decimal(0)
                product = Product(code=code, **fields, slug=product_slug(designation, brand, code),
                    search_key=normalize_designation(designation), has_spaces=any(c.isspace() for c in designation.strip()))
                new_products[code] = product
            warehouse_name = row[10] or DEFAULT_WAREHOUSE
            stock_specs.append((code, warehouses[warehouse_name], int(row[11] or 0), int(row[12] or 0)))
            price_specs[code] = (row[14], row[15])
        Product.objects.bulk_create(new_products.values(), batch_size=500)
        products = {p.code: p for p in Product.objects.filter(code__in=[s[0] for s in stock_specs])}
        new_stocks, new_prices = [], []
        for code, warehouse, amount, reserved in stock_specs:
            if code in existing:
                Stock.objects.update_or_create(product=products[code], warehouse=warehouse,
                    defaults={'amount': amount, 'reserved': reserved, 'free': amount - reserved})
            else:
                new_stocks.append(Stock(product=products[code], warehouse=warehouse, amount=amount, reserved=reserved, free=amount-reserved))
        for code, (retail, wholesale) in price_specs.items():
            for kind, value in [('retail', retail), ('wholesale', wholesale)]:
                if code in existing:
                    if value is not None:
                        Price.objects.update_or_create(product=products[code], price_type=kind, defaults={'price': value})
                elif value is not None:
                    new_prices.append(Price(product=products[code], price_type=kind, price=value))
        Stock.objects.bulk_create(new_stocks, batch_size=500)
        Price.objects.bulk_create(new_prices, batch_size=500)
    return result
