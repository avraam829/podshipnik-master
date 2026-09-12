from io import BytesIO
from tempfile import NamedTemporaryFile
from unittest.mock import Mock
import hashlib

from django.contrib import admin
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import RequestFactory, TestCase
from openpyxl import Workbook, load_workbook

from catalog.admin import ProductAdmin
from catalog.importing import EXCEL_HEADERS, ImportValidationError, import_xlsx
from catalog.models import Price, Product, Stock, Warehouse


def row(code='', designation='6204', warehouse='Склад А', amount=10, reserved=2,
        retail=100, wholesale=80):
    return [
        'Подшипники', 'Шариковые', code, designation, 'SKF', 'Описание',
        0.2, 20, 47, 14, warehouse, amount, reserved, amount - reserved,
        retail, wholesale,
    ]


class ExcelImportWorkflowTests(TestCase):
    def xlsx(self, rows, headers=EXCEL_HEADERS):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for values in rows:
            sheet.append(values)
        temporary = NamedTemporaryFile(suffix='.xlsx', delete=False)
        temporary.close()
        workbook.save(temporary.name)
        self.addCleanup(lambda: __import__('os').unlink(temporary.name))
        return temporary.name

    def test_preview_is_read_only_and_apply_creates_product_with_multiple_stocks(self):
        path = self.xlsx([
            row('SKU-6204', warehouse='Склад А'),
            row('SKU-6204', warehouse='Склад Б', amount=4, reserved=1),
        ])

        preview = import_xlsx(path)
        self.assertEqual(preview['created'], 1)
        self.assertEqual(Product.objects.count(), 0)

        result = import_xlsx(path, dry_run=False)
        self.assertEqual(result['created'], 1)
        product = Product.objects.get(code='SKU-6204')
        self.assertEqual(Stock.objects.filter(product=product).count(), 2)
        self.assertEqual(Price.objects.filter(product=product).count(), 2)

    def test_empty_price_does_not_erase_existing_price(self):
        first = self.xlsx([row('SKU-6204', retail=100, wholesale=80)])
        import_xlsx(first, dry_run=False)
        changed = self.xlsx([row('SKU-6204', designation='6204-2', retail=None, wholesale=None)])

        import_xlsx(changed, dry_run=False)

        product = Product.objects.get(code='SKU-6204')
        self.assertEqual(product.designation, '6204-2')
        self.assertEqual(product.price, 100)
        self.assertEqual(
            dict(Price.objects.filter(product=product).values_list('price_type', 'price')),
            {'retail': 100, 'wholesale': 80},
        )

    def test_changed_no_code_file_is_rejected_but_exact_repeat_is_a_noop(self):
        original = self.xlsx([row(code='', designation='6305')])
        import_xlsx(original, dry_run=False)
        exact_repeat = import_xlsx(original, dry_run=False)
        self.assertEqual(exact_repeat['unchanged'], 1)
        self.assertEqual(Product.objects.count(), 1)

        changed = self.xlsx([row(code='', designation='6305-изменён')])
        with self.assertRaisesRegex(ImportValidationError, 'без кода'):
            import_xlsx(changed, dry_run=False)
        self.assertEqual(Product.objects.count(), 1)

    def test_price_error_is_a_warning_and_uses_existing_price(self):
        initial = self.xlsx([row('SKU-6204', retail=100)])
        import_xlsx(initial, dry_run=False)
        broken = self.xlsx([row('SKU-6204', retail='#VALUE!')])

        result = import_xlsx(broken, dry_run=False)

        self.assertEqual(len(result['warnings']), 1)
        self.assertEqual(Product.objects.get(code='SKU-6204').price, 100)

    def test_unique_legacy_warehouse_name_is_reused_without_duplicate_stock(self):
        warehouse = Warehouse.objects.create(code='LEGACY-A', name='Склад А')
        product = Product.objects.create(code='SKU-6204', designation='6204')
        Stock.objects.create(product=product, warehouse=warehouse, amount=1, reserved=0)

        import_xlsx(self.xlsx([row('SKU-6204', amount=10, reserved=2)]), dry_run=False)

        self.assertEqual(Warehouse.objects.count(), 1)
        stock = Stock.objects.get(product=product, warehouse=warehouse)
        self.assertEqual((stock.amount, stock.reserved, stock.free), (10, 2, 8))

    def test_ambiguous_warehouse_name_is_rejected_before_apply(self):
        Warehouse.objects.create(code='LEGACY-A1', name='Склад А')
        Warehouse.objects.create(code='LEGACY-A2', name='Склад А')

        with self.assertRaisesRegex(ImportValidationError, 'несколько раз'):
            import_xlsx(self.xlsx([row('SKU-6204')]), dry_run=False)
        self.assertEqual(Product.objects.count(), 0)

    def test_hashed_warehouse_code_collision_is_rejected_in_preview(self):
        name = 'Склад А'
        code = 'WH-' + hashlib.sha256(name.encode()).hexdigest()[:20]
        Warehouse.objects.create(code=code, name='Другой склад')

        with self.assertRaisesRegex(ImportValidationError, 'уже занят'):
            import_xlsx(self.xlsx([row('SKU-6204', warehouse=name)]))
        self.assertEqual(Product.objects.count(), 0)

    def test_template_endpoint_exports_accepted_headers(self):
        request = RequestFactory().get('/admin/catalog/product/import-template/')
        request.user = Mock()
        request.user.has_perm.return_value = True

        response = ProductAdmin(Product, admin.site).import_template(request)

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(tuple(next(workbook.active.iter_rows(values_only=True))), EXCEL_HEADERS)

    def test_legacy_json_command_is_disabled_without_writing(self):
        with self.assertRaises(CommandError):
            call_command('import_catalog')
        self.assertEqual(Product.objects.count(), 0)
